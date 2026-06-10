import os
import io
import re
import time
import requests

API_BASE = "http://127.0.0.1:8000"
PDF_FILES = {
    r"D:\OCR_MASTER\datasets\salesinvoice.pdf": 85,
    r"D:\OCR_MASTER\datasets\Glsprint.pdf": 148,
    r"D:\OCR_MASTER\datasets\POPRINT.pdf": 156,
    r"D:\OCR_MASTER\datasets\GRN.pdf": 62,
}
REQUIRED_PDF_FIELDS = {
    "salesinvoice.pdf": ("invoice_number", "invoice_date", "buyer_name", "grand_total"),
    "Glsprint.pdf": ("invoice_number", "invoice_date", "consignee", "grand_total"),
    "POPRINT.pdf": ("po_number", "po_date"),
    "GRN.pdf": ("grn_number", "grn_date"),
}
REQUIRED_XLSX_FIELDS = REQUIRED_PDF_FIELDS
FIELD_ALIASES = {
    "Glsprint.pdf": {
        "consignee": ("consignee", "consignee_name", "buyer_name", "customer_name", "party_name"),
        "grand_total": ("grand_total", "total_amount", "invoice_total", "net_total", "total_value"),
    },
    "POPRINT.pdf": {
        "po_date": ("po_date", "purchase_order_date", "order_date", "document_date", "invoice_date", "date"),
    },
}


def _field_value(fields, key):
    value = fields.get(key)
    if isinstance(value, dict):
        value = value.get("value")
    if value in (None, "", [], {}):
        return ""
    return str(value).strip()


def _field_value_with_aliases(name, fields, key):
    for candidate in FIELD_ALIASES.get(name, {}).get(key, (key,)):
        value = _field_value(fields, candidate)
        if value:
            return value, candidate
    return "", key


def _normalize_text(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _extract_pdf_text(pdf_bytes):
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(pdf_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        # Fallback for simple/uncompressed PDFs; keeps the test dependency-free.
        decoded = pdf_bytes.decode("latin-1", errors="ignore")
        return "".join(ch if ch.isprintable() or ch.isspace() else " " for ch in decoded)


def _assert_pdf_content(name, fields, pdf_bytes):
    required = REQUIRED_PDF_FIELDS.get(name, ())
    resolved = {key: _field_value_with_aliases(name, fields, key) for key in required}
    missing = [key for key, (value, _) in resolved.items() if not value]
    if missing:
        return False, f"missing required fields: {', '.join(missing)}"

    pdf_text = _normalize_text(_extract_pdf_text(pdf_bytes))
    absent = []
    for key, (raw_value, _) in resolved.items():
        value = _normalize_text(raw_value)
        if value and value not in pdf_text:
            absent.append(key)
    if absent:
        return False, f"field values absent from PDF text: {', '.join(absent)}"
    details = ", ".join(
        key if alias == key else f"{key}->{alias}"
        for key, (_, alias) in resolved.items()
    )
    return True, f"required fields present in PDF: {details}"


def _extract_xlsx_text(xlsx_bytes):
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
        values = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value not in (None, ""):
                        values.append(str(value))
        workbook.close()
        return "\n".join(values)
    except Exception as exc:
        return f"__XLSX_PARSE_ERROR__ {exc}"


def _assert_xlsx_content(name, fields, xlsx_bytes):
    required = REQUIRED_XLSX_FIELDS.get(name, ())
    resolved = {key: _field_value_with_aliases(name, fields, key) for key in required}
    missing = [key for key, (value, _) in resolved.items() if not value]
    if missing:
        return False, f"missing required fields: {', '.join(missing)}"

    workbook_text = _normalize_text(_extract_xlsx_text(xlsx_bytes))
    if "__xlsx_parse_error__" in workbook_text:
        return False, "workbook parse failed"

    absent = []
    for key, (raw_value, _) in resolved.items():
        value = _normalize_text(raw_value)
        if value and value not in workbook_text:
            absent.append(key)
    if absent:
        return False, f"field values absent from workbook: {', '.join(absent)}"
    details = ", ".join(
        key if alias == key else f"{key}->{alias}"
        for key, (_, alias) in resolved.items()
    )
    return True, f"required fields present in workbook: {details}"

def main():
    print("Testing Backend startup...")
    try:
        res = requests.get(f"{API_BASE}/docs")
        if res.status_code == 200:
            print("Backend startup: PASS")
        else:
            print(f"Backend startup: FAIL (status code {res.status_code})")
            print("FAILED")
            print("ROOT CAUSE: Backend startup check failed")
            print("FILES TO PATCH: app/main.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix backend startup")
            return
    except Exception as e:
        print(f"Backend startup: FAIL ({e})")
        print("FAILED")
        print("ROOT CAUSE: Backend is not reachable at http://127.0.0.1:8000")
        print("FILES TO PATCH: None")
        print("COMMANDS: uvicorn app.main:app --host 127.0.0.1 --port 8000")
        print("NEXT TASK: Start backend and rerun")
        return

    print("Testing Login...")
    # register a test user
    email = "test@example.com"
    password = "password123"
    try:
        requests.post(f"{API_BASE}/api/v1/auth/register", json={
            "email": email, "password": password, "full_name": "Test User"
        })
    except:
        pass
    
    # login
    res = requests.post(f"{API_BASE}/api/v1/auth/login", json={
        "email": email, "password": password
    })
    if res.status_code != 200:
        print(f"Login: FAIL (status {res.status_code} - {res.text})")
        print("FAILED")
        print("ROOT CAUSE: Login failed")
        print("FILES TO PATCH: app/api/v1/routers/auth.py")
        print("COMMANDS: python scripts/e2e_test.py")
        print("NEXT TASK: Fix auth")
        return
    token = res.json()["access_token"]
    print("Login: PASS")

    headers = {"Authorization": f"Bearer {token}"}

    results = []
    for pdf, expected_count in PDF_FILES.items():
        name = os.path.basename(pdf)
        print(f"\nProcessing {name}...")
        if not os.path.exists(pdf):
            print("FAILED")
            print(f"ROOT CAUSE: Dataset not found {pdf}")
            print("FILES TO PATCH: None")
            print("COMMANDS: None")
            print("NEXT TASK: Restore dataset file")
            return
        
        # Upload
        with open(pdf, "rb") as f:
            res = requests.post(f"{API_BASE}/api/v1/upload/", headers=headers, files={"file": f})
        
        if res.status_code != 202:
            print(f"  Upload FAIL: {res.text}")
            print("FAILED")
            print(f"ROOT CAUSE: Upload failed for {name}")
            print("FILES TO PATCH: app/api/v1/routers/upload.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix upload path")
            return
        
        job_id = res.json()["job_id"]
        
        # Poll
        job_data = None
        for _ in range(60):
            time.sleep(2)
            res = requests.get(f"{API_BASE}/api/v1/upload/jobs/{job_id}", headers=headers)
            if res.status_code == 200:
                data = res.json()
                if data["status"] == "completed":
                    print("  OCR: completed")
                    job_data = data
                    break
                elif data["status"] == "failed":
                    print(f"  OCR: failed ({data.get('error_message')})")
                    break
            else:
                print(f"  Poll FAIL: {res.text}")
                break
        else:
            print("  OCR timeout")
        
        if not job_data or not job_data.get("extracted_data_id"):
            print("FAILED")
            print(f"ROOT CAUSE: OCR job did not produce extracted_data_id for {name}")
            print("FILES TO PATCH: app/api/v1/routers/upload.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix OCR job completion")
            return

        raw_text = (
            job_data.get("raw_text")
            or job_data.get("text")
            or job_data.get("full_text")
            or ""
        ).lower()
        mock_signals = ("mock ocr", "install paddlepaddle", "no real ocr engine")
        if (
            job_data.get("is_mock_ocr")
            or job_data.get("ocr_engine") == "mock"
            or any(signal in raw_text for signal in mock_signals)
        ):
            print("FAILED")
            print(f"ROOT CAUSE: Mock OCR used for {name}")
            print("FILES TO PATCH: app/services/ocr_service.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Restore real OCR engine")
            return

        errors = job_data.get("validation_errors") or []
        if errors:
            print("FAILED")
            print(f"ROOT CAUSE: Validation errors before export for {name}: {errors}")
            print("FILES TO PATCH: app/services/validation_service.py, app/api/v1/routers/upload.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validation failures")
            return

        extraction_status = job_data.get("extraction_status")
        if extraction_status != "validated":
            print("FAILED")
            print(f"ROOT CAUSE: Extraction status is {extraction_status!r}, expected 'validated' for {name}")
            print("FILES TO PATCH: app/api/v1/routers/upload.py, app/services/validation_service.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validation status")
            return

        fields = job_data.get("mapped_fields") or {}
        actual_count = len(fields)
        print(f"  Validated fields: {actual_count}/{expected_count}")
        if actual_count < expected_count:
            print("FAILED")
            print(f"ROOT CAUSE: Field count drift for {name}: expected >= {expected_count}, got {actual_count}")
            print("FILES TO PATCH: app/services/company_form_mapper.py, app/services/validation_service.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Preserve verified field counts")
            return

        extracted_data_id = job_data["extracted_data_id"]
        payload = {"fields": fields, "extracted_data_id": extracted_data_id}

        # Export PDF
        res = requests.post(f"{API_BASE}/api/v1/forms/export/pdf", headers=headers, json=payload)
        if res.status_code != 200:
            print(f"  PDF Export FAIL: {res.text}")
            print("FAILED")
            print(f"ROOT CAUSE: PDF export failed for {name}")
            print("FILES TO PATCH: app/api/v1/routers/forms.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validated PDF export")
            return
        if not res.content.startswith(b"%PDF"):
            print("  PDF Export FAIL: response does not start with %PDF")
            print("FAILED")
            print(f"ROOT CAUSE: PDF export returned non-PDF content for {name}")
            print("FILES TO PATCH: app/api/v1/routers/forms.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validated PDF export")
            return
        if len(res.content) <= 100:
            print(f"  PDF Export FAIL: PDF too small ({len(res.content)} bytes)")
            print("FAILED")
            print(f"ROOT CAUSE: PDF export returned undersized PDF for {name}")
            print("FILES TO PATCH: app/api/v1/routers/forms.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validated PDF export")
            return

        pdf_ok, pdf_detail = _assert_pdf_content(name, fields, res.content)
        if not pdf_ok:
            print(f"  PDF Export FAIL: {pdf_detail}")
            print("FAILED")
            print(f"ROOT CAUSE: PDF content assertion failed for {name}: {pdf_detail}")
            print("FILES TO PATCH: app/api/v1/routers/forms.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validated PDF export content")
            return
        print(f"  PDF Export: PASS ({pdf_detail})")

        # Export Excel
        res = requests.post(f"{API_BASE}/api/v1/forms/export/excel", headers=headers, json=payload)
        if res.status_code != 200:
            print(f"  Excel Export FAIL: {res.text}")
            print("FAILED")
            print(f"ROOT CAUSE: Excel export failed for {name}")
            print("FILES TO PATCH: app/api/v1/routers/forms.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validated Excel export")
            return
        if not res.content.startswith(b"PK"):
            print("  Excel Export FAIL: response is not a valid XLSX zip")
            print("FAILED")
            print(f"ROOT CAUSE: Excel export returned non-XLSX content for {name}")
            print("FILES TO PATCH: app/api/v1/routers/forms.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validated Excel export")
            return
        if len(res.content) <= 100:
            print(f"  Excel Export FAIL: workbook too small ({len(res.content)} bytes)")
            print("FAILED")
            print(f"ROOT CAUSE: Excel export returned undersized workbook for {name}")
            print("FILES TO PATCH: app/api/v1/routers/forms.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validated Excel export")
            return

        xlsx_ok, xlsx_detail = _assert_xlsx_content(name, fields, res.content)
        if not xlsx_ok:
            print(f"  Excel Export FAIL: {xlsx_detail}")
            print("FAILED")
            print(f"ROOT CAUSE: Excel content assertion failed for {name}: {xlsx_detail}")
            print("FILES TO PATCH: app/api/v1/routers/forms.py")
            print("COMMANDS: python scripts/e2e_test.py")
            print("NEXT TASK: Fix validated Excel export content")
            return
        print(f"  Excel Export: PASS ({xlsx_detail})")

        results.append((name, actual_count))

    print("\nCOMPLETED")
    for name, count in results:
        print(f"{name}: {count}")
    print("ROOT CAUSE: None")
    print("FILES TO PATCH: None")
    print("COMMANDS: python scripts/e2e_test.py")
    print("NEXT TASK: None")

if __name__ == '__main__':
    main()
