from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MAX_TABULAR_ROWS = 5000
MAX_TABULAR_SHEETS = 10


def _key(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[\s_\-/]+", " ", text)
    return re.sub(r"[^a-z0-9 ]+", "", text).strip()


FIELD_ALIASES = {
    "sales_order_no": {"sales order no", "sales order number", "so no", "so number", "customer po no", "customer po number"},
    "sales_order_date": {"sales order date", "so date"},
    "customer_name": {"customer", "buyer", "bill to", "customer name", "party name"},
    "consignee_name": {"consignee", "ship to", "delivery to"},
    "po_number": {"purchase order no", "purchase order number", "po no", "po number"},
    "po_date": {"po date", "purchase order date"},
    "supplier_name": {"supplier", "vendor", "supplier name", "vendor name"},
    "payment_terms": {"payment term", "payment terms", "terms"},
    "delivery_terms": {"delivery term", "delivery terms"},
    "currency": {"currency", "currency code"},
    "remarks": {"remarks", "notes"},
}

ROW_ALIASES = {
    "product": {"product", "item", "item name", "material", "article"},
    "part_no": {"part no", "part number", "item code", "product code", "sku"},
    "description": {"description", "product description", "item description"},
    "hsn_code": {"hsn", "hsn code", "hsn sac"},
    "qty": {"qty", "quantity", "order qty", "received qty"},
    "uom": {"uom", "unit", "units"},
    "unit_price": {"price", "rate", "unit price"},
    "discount_amount": {"discount", "disc amount"},
    "tax_amount": {"tax", "gst amount", "igst amount"},
    "tax_group": {"tax group", "gst rate", "tax rate"},
    "line_total": {"total", "amount", "line amount", "net amount"},
}

FIELD_LOOKUP = {alias: name for name, aliases in FIELD_ALIASES.items() for alias in aliases}
ROW_LOOKUP = {alias: name for name, aliases in ROW_ALIASES.items() for alias in aliases}
GENERIC_DATE_ALIASES = {"date", "order date"}
GENERIC_ORDER_ALIASES = {"order no", "order number"}


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _canonical_field(label: Any, context: set[str] | None = None) -> str | None:
    normalized = _key(label)
    if normalized in FIELD_LOOKUP:
        return FIELD_LOOKUP[normalized]
    context = context or set()
    purchase_context = bool(context & {"po_number", "po_date", "supplier_name"})
    if normalized in GENERIC_DATE_ALIASES:
        return "po_date" if purchase_context else "sales_order_date"
    if normalized in GENERIC_ORDER_ALIASES:
        return "po_number" if purchase_context else "sales_order_no"
    return None


def _header_map(row: list[Any]) -> dict[int, str]:
    result: dict[int, str] = {}
    for index, value in enumerate(row):
        canonical = ROW_LOOKUP.get(_key(value))
        if canonical:
            result[index] = canonical
    return result


def _find_header(rows: list[list[Any]]) -> tuple[int | None, dict[int, str]]:
    best_index: int | None = None
    best_map: dict[int, str] = {}
    for index, row in enumerate(rows[:100]):
        mapping = _header_map(row)
        if len(mapping) >= 2 and len(mapping) > len(best_map):
            best_index, best_map = index, mapping
    return best_index, best_map


def _extract_key_values(rows: list[list[Any]], detected_fields: dict[str, Any]) -> dict[str, Any]:
    mapped: dict[str, Any] = {}
    context: set[str] = set()
    for row in rows[:250]:
        for cell in row:
            canonical = _canonical_field(cell, context)
            if canonical:
                context.add(canonical)
    for row in rows[:250]:
        for index, label in enumerate(row[:-1]):
            canonical = _canonical_field(label, context)
            if not canonical:
                continue
            value = next((_display(item) for item in row[index + 1:] if _display(item)), "")
            if not value or canonical in mapped:
                continue
            original_label = _display(label)
            meta = {"value": value, "confidence": 0.95, "method": "tabular_key_value", "source_label": original_label}
            mapped[canonical] = meta
            detected_fields[original_label] = meta.copy()
    return mapped


def _extract_table(rows: list[list[Any]], sheet_name: str = "") -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    header_index, mapping = _find_header(rows)
    if header_index is None:
        return [], None
    original_headers = [_display(value) for value in rows[header_index]]
    item_rows: list[dict[str, Any]] = []
    original_rows: list[dict[str, str]] = []
    for row in rows[header_index + 1:]:
        item: dict[str, Any] = {}
        original: dict[str, str] = {}
        for column, canonical in mapping.items():
            value = _display(row[column] if column < len(row) else "")
            if value:
                item[canonical] = value
                original[original_headers[column] or canonical] = value
        if item:
            if sheet_name:
                item["sheet_name"] = sheet_name
            item_rows.append(item)
            original_rows.append(original)
    table = {"header_row": header_index + 1, "headers": original_headers, "original_rows": original_rows}
    return item_rows, table


def _classify(mapped: dict[str, Any], raw_text: str) -> tuple[str, str]:
    keys = set(mapped)
    text = _key(raw_text)
    if any(term in text for term in ("grn", "goods receipt", "received qty")):
        return "unknown", "GRN"
    purchase_score = 2 * len(keys & {"supplier_name", "po_number", "po_date"})
    sales_score = 2 * len(keys & {"customer_name", "sales_order_no", "sales_order_date", "consignee_name"})
    purchase_score += sum(term in text for term in ("purchase order", "supplier", "vendor"))
    sales_score += sum(term in text for term in ("sales order", "customer", "buyer", "customer po"))
    if purchase_score > sales_score and purchase_score:
        return "purchase_order", "PURCHASE_ORDER"
    if sales_score > purchase_score and sales_score:
        return "sales_order", "SALES_ORDER"
    return "unknown", "GENERAL_TABULAR"


def normalize_tabular_document(parsed_sheets: list[dict[str, Any]], source_type: str = "excel", warnings: list[str] | None = None) -> dict[str, Any]:
    warnings = list(warnings or [])
    mapped_fields: dict[str, Any] = {}
    detected_fields: dict[str, Any] = {}
    item_rows: list[dict[str, Any]] = []
    raw_lines: list[str] = []
    sheets: list[dict[str, Any]] = []
    remaining_rows = MAX_TABULAR_ROWS

    for sheet in parsed_sheets[:MAX_TABULAR_SHEETS]:
        rows = [list(row) for row in sheet.get("rows", []) if any(_display(value) for value in row)]
        if not rows:
            continue
        if len(rows) > remaining_rows:
            rows = rows[:remaining_rows]
            warnings.append(f"Row limit of {MAX_TABULAR_ROWS} reached; remaining rows were truncated.")
        remaining_rows -= len(rows)
        name = str(sheet.get("name") or "Sheet")
        raw_lines.append(f"[{name}]")
        raw_lines.extend("\t".join(_display(value) for value in row).rstrip() for row in rows)
        local_detected: dict[str, Any] = {}
        for key, value in _extract_key_values(rows, local_detected).items():
            mapped_fields.setdefault(key, value)
        detected_fields.update({key: value for key, value in local_detected.items() if key not in detected_fields})
        extracted_rows, table = _extract_table(rows, name if source_type == "excel" else "")
        item_rows.extend(extracted_rows)
        sheets.append({"name": name, "row_count": len(rows), "table": table})
        if remaining_rows <= 0:
            break

    if len(parsed_sheets) > MAX_TABULAR_SHEETS:
        warnings.append(f"Sheet limit of {MAX_TABULAR_SHEETS} reached; remaining sheets were skipped.")
    if not sheets:
        raise ValueError("The Excel/CSV document is empty or contains no readable values.")
    if item_rows:
        mapped_fields["item_rows"] = {"value": item_rows, "confidence": 0.95, "method": "tabular_table"}
    raw_text = "\n".join(line for line in raw_lines if line).strip()
    module, document_type = _classify(mapped_fields, raw_text)
    mapped_fields["document_type"] = {"value": document_type, "confidence": 0.95, "method": "tabular_classifier"}
    return {
        "raw_text": raw_text,
        "extracted_text": raw_text,
        "full_text": raw_text,
        "mapped_fields": mapped_fields,
        "canonical_fields": mapped_fields.copy(),
        "detected_fields": detected_fields,
        "item_rows": item_rows,
        "items": item_rows,
        "sheets": sheets,
        "document_type": document_type,
        "module": module,
        "warnings": list(dict.fromkeys(warnings)),
        "source_type": source_type,
        "detected_count": len([key for key in mapped_fields if key not in {"item_rows", "document_type"}]) + sum(len(row) for row in item_rows),
    }


def parse_excel_document(file_bytes: bytes, filename: str) -> dict[str, Any]:
    if Path(filename or "").suffix.lower() == ".xls":
        raise ValueError("Legacy .xls is not supported. Please upload .xlsx or .csv.")
    if not file_bytes:
        raise ValueError("The Excel document is empty.")
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Invalid Excel workbook: {exc}") from exc
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        parsed = [{"name": sheet.title, "rows": [list(row) for row in sheet.iter_rows(values_only=True)]} for sheet in visible[:MAX_TABULAR_SHEETS]]
        warnings = []
        if len(visible) > MAX_TABULAR_SHEETS:
            warnings.append(f"Sheet limit of {MAX_TABULAR_SHEETS} reached; remaining sheets were skipped.")
        return normalize_tabular_document(parsed, "excel", warnings)
    finally:
        workbook.close()


def parse_csv_document(file_bytes: bytes, filename: str) -> dict[str, Any]:
    if not file_bytes:
        raise ValueError("The CSV document is empty.")
    text = ""
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text.strip():
        raise ValueError("The CSV document is empty.")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = max(",;\t|", key=lambda candidate: sample.count(candidate))
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    result = normalize_tabular_document([{"name": Path(filename or "document.csv").name, "rows": rows}], "csv")
    result["delimiter"] = delimiter
    return result
